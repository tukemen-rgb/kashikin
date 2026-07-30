# -*- coding: utf-8 -*-
"""Claude Design は XLSX を直接読めるので、見本データを表としても書き出す。"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

S = json.load(open('design/data/sample-questions.json'))
Q, M = S['questions'], S['meta']
FONT = 'Meiryo'                      # アプリ本体と同じ書体に合わせる
HEAD = PatternFill('solid', fgColor='0A5A2C')      # デザインシステムのヘッダー色
BAND = PatternFill('solid', fgColor='E5F1E9')
THIN = Border(*[Side('thin', color='D4D8DE')] * 4)
wb = Workbook()

def style(ws, widths, freeze='A2'):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        c.fill = HEAD
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = THIN
    ws.freeze_panes = freeze

# ---------------------------------------------------------------- 設問
ws = wb.active
ws.title = '設問'
COLS = ['ID', '回', '年度', '問', '分野', 'テーマ', '形式', '否定形', '設問',
        '記述ａ', '記述ｂ', '記述ｃ', '記述ｄ',
        '選択肢①', '選択肢②', '選択肢③', '選択肢④',
        '注記', '正答', '条文', '引っかけの型', '頻出度',
        '見本としての役割', '設問字数', '記述の最長', '選択肢の最長']
ws.append(COLS)
FMT = {'single': '単純択一', 'combo': '組合せ', 'count': '個数'}
MARK = ['①', '②', '③', '④']
for q in Q:
    st = q['statements']
    op = q['options']
    ws.append([q['id'], q['label'], q['era'], q['num'], q['section'], q['theme'],
               FMT[q['format']], '○' if q['negative'] else '',
               q['stem'],
               st.get('ａ', ''), st.get('ｂ', ''), st.get('ｃ', ''), st.get('ｄ', ''),
               op.get('①', ''), op.get('②', ''), op.get('③', ''), op.get('④', ''),
               q['note'] or '', MARK[q['answer'] - 1],
               '\n'.join(q['articles']), '、'.join(q['traps']), q['topicRank'],
               q['sampleFor'], q['len']['stem'], q['len']['maxStatement'], q['len']['maxOption']])
style(ws, [9, 8, 11, 5, 15, 24, 10, 7, 62, 40, 40, 40, 40, 40, 40, 40, 40,
           30, 6, 30, 22, 7, 30, 9, 9, 10], freeze='B2')
for r in range(2, len(Q) + 2):
    ws.row_dimensions[r].height = 78
    if r % 2 == 0:
        for c in ws[r]:
            c.fill = BAND

# ---------------------------------------------------------------- 引っかけ
ws = wb.create_sheet('引っかけの型')
ws.append(['型ID', '表示名', '着眼点', '見本での出現数'])
for tid, t in M['引っかけの型'].items():
    ws.append([tid, t['表示名'], t['着眼点'],
               sum(1 for q in Q if tid in q['traps'])])
style(ws, [12, 20, 78, 15])
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 46

# ---------------------------------------------------------------- 目安
ws = wb.create_sheet('レイアウトの目安')
ws.append(['項目', '最短', '中央', '最長', '備考'])
c = M['文字数の実測']
ws.append(['設問', c['設問']['最短'], c['設問']['中央'], c['設問']['最長'],
           '最長は第1回 問4。括弧の外の「。」で9文に割れる'])
ws.append(['選択肢1つ', c['選択肢']['最短'], c['選択肢']['中央'], c['選択肢']['最長'],
           '最長は第3回 問47'])
ws.append(['記述ａ〜ｄ 1つ', '', '', c['記述ａ〜ｄ']['最長'], '最長は第6回 問1'])
ws.append([])
ws.append(['分野', '問題数', '見本の件数', '', ''])
row0 = ws.max_row + 1
for sec, n in M['分野の内訳'].items():
    ws.append([sec, n, sum(1 for q in Q if q['section'] == sec), '', ''])
ws.append(['合計', sum(M['分野の内訳'].values()), len(Q), '', ''])
ws.append([])
ws.append(['出題形式', '見本の件数', '', '', ''])
for k, v in M['出題形式'].items():
    ws.append([v, sum(1 for q in Q if q['format'] == k), '', '', ''])
style(ws, [20, 10, 12, 10, 52])
ws['A6'].font = Font(name=FONT, bold=True, size=10)

# ---------------------------------------------------------------- 読み方
ws = wb.create_sheet('はじめに')
NOTE = [
    ['貸金業務取扱主任者 過去問 — デザイン用の見本データ', ''],
    ['', ''],
    ['この表は何か',
     '画面設計に使うための見本です。全1000問のうち、レイアウトが壊れる条件を'
     '網羅するように24問を選んであります（最長・最短、3つの出題形式、'
     '4分野、注記つき、図表つき、引っかけの有無）。'],
    ['守ること',
     '設問・記述・選択肢・注記は公式PDFの原文そのままです。'
     '表示上の改行や強調は加えてかまいませんが、'
     '文字を足す・削る・置き換えることはできません。'],
    ['全体の規模',
     f"{M['全体']['問題数']}問／{M['全体']['回']}回分／"
     f"{M['全体']['分野']}分野／{M['全体']['テーマ']}テーマ／"
     f"引っかけの型 {M['全体']['引っかけの型']}種"],
    ['シート', '設問：見本24問　／　引っかけの型：12種と着眼点　／　'
               'レイアウトの目安：字数の実測と件数'],
    ['全1000問',
     'data/questions.json（プログラム用）と data/questions.csv（表計算用）にあります。'],
    ['配色',
     '深緑 #0F7A3C＝操作　／　朱 #C6402A＝先へ進む・誤り　／　'
     '紺 #134E8C＝できた　／　紫 #6A4A9C＝あとで見直す。'
     'コンポーネントは design/ に12枚のカードとして置いてあります。'],
]
for a, b in NOTE:
    ws.append([a, b])
ws.column_dimensions['A'].width = 16
ws.column_dimensions['B'].width = 96
ws['A1'].font = Font(name=FONT, bold=True, size=14, color='0A5A2C')
for r in range(3, len(NOTE) + 1):
    ws[f'A{r}'].font = Font(name=FONT, bold=True, size=10)
    ws[f'B{r}'].font = Font(name=FONT, size=10)
    for col in 'AB':
        ws[f'{col}{r}'].alignment = Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[r].height = 44
wb.move_sheet('はじめに', -3)

wb.save('design/data/sample-questions.xlsx')
print('design/data/sample-questions.xlsx')
